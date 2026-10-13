"""Coverage-matrix builder and guard: hand content survives, statuses parse."""
from openpyxl import Workbook

from conf_abstracts import build_coverage_matrix as B
from conf_abstracts import check_coverage_matrix as G


def test_matrix_status_parse_location_with_commas():
    assert G.parse_status("Vancouver, Canada; Missing — 12 abstracts", "L") == "Missing"
    assert G.parse_status("?; Missing — 1 queue row", "Z") == "Missing"
    assert G.parse_status("?; Missing — 1 queue row", "AA") == "Missing"
    assert G.parse_status("Missing — 5 abstracts cited", "Y") == "Missing"
    assert G.parse_status("Guelph, ON; Ingested", "B") == "Ingested"
    assert "Partial" in G.STATUSES


def test_matrix_manual_store_has_the_hand_series():
    heads = [c["heading"] for c in B.MANUAL_COLS]
    assert heads == ["SVP", "SICB", "SIBM"]
    assert [r["series"][:4] for r in B.MANUAL_CONFS] == ["SVP ", "SICB", "SIBM"]


def test_matrix_cell_note_prepends_and_retires():
    wb = Workbook(); ws = wb.active
    ws.append(["Year", "ASIH/JMIH", "SOMEPEC"])
    ws.append([1998, "Guelph, ON; Ingested", "x; Ingested — 44 abstracts"])
    ws.append([2001, "University Park, PA; OCR — degraded", ""])
    saved = B.MANUAL_NOTES
    B.MANUAL_NOTES = [
        dict(column="ASIH/JMIH", year=1998, status="Ingested", note="Fable-merged 2026-09-17"),
        dict(column="SOMEPEC", year=1998, status="Ingested", note="Fable extraction 2026-09-17"),
        dict(column="ASIH/JMIH", year=2001, status="Ingested", note="flatbed rescan")]
    try:
        B._apply_cell_notes(ws, ["Year", "ASIH/JMIH", "SOMEPEC"])
    finally:
        B.MANUAL_NOTES = saved
    assert ws["B2"].value == "Guelph, ON; Ingested — Fable-merged 2026-09-17"
    assert ws["C2"].value == "x; Ingested — Fable extraction 2026-09-17, 44 abstracts"
    assert ws["B3"].value == "University Park, PA; OCR — degraded"     # status moved on: retired

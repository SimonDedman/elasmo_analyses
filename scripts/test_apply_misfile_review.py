#!/usr/bin/env python3
"""Tests for the misfile applier.

The distinction this has to hold onto is between deleting a NAME and
deleting CONTENT. Five filenames can be five links to one document, and the
count that matters is per inode, never per path.
"""

import os
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))

import apply_misfile_review as app
from dedupe_hardlink import sha256

BODY = b"%PDF-1.4 an annotated bibliography of parasitic isopoda" + b"x" * 2048

COLS = ["group_id", "how_sure", "record_year", "absent_title", "pdf_year",
        "pdf_holds_instead", "correct_copy_elsewhere", "open_correct",
        "pdf_starts", "english_gloss", "title_words_found", "open_pdf",
        "decision", "notes", "why", "n_records", "size_mb", "words_found",
        "words_sought", "latin_fraction", "n_words", "absent_path",
        "correct_copy_path", "sha256"]


def workbook(tmp_path, entries):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "misfiled"
    ws.append(COLS)
    for path, decision, correct in entries:
        row = {c: "" for c in COLS}
        row.update({"group_id": 1, "how_sure": "strong", "record_year": 2004,
                    "absent_title": "Birth of guitarfish",
                    "pdf_holds_instead": "Parasitic Isopoda bibliography",
                    "decision": decision, "absent_path": str(path),
                    "correct_copy_path": correct,
                    "sha256": sha256(Path(path))[:12]})
        ws.append([row[c] for c in COLS])
    out = tmp_path / "review.xlsx"
    wb.save(out)
    return out


@pytest.fixture
def linked_pair(tmp_path):
    """One document, two names: the correct one and a misfiled one."""
    good = tmp_path / "Moreira.1978.Parasitic Isopoda bibliography.pdf"
    good.write_bytes(BODY)
    bad = tmp_path / "Gonzalez.2004.Birth.of.guitarfish.pdf"
    os.link(good, bad)
    return good, bad


def run(book, apply=True, allow_loss=False):
    actions, _ = app.plan(app.read_rows(book), log=lambda *_: None)
    freed, emptied = app.execute(actions, apply=apply, allow_loss=allow_loss,
                                 log=lambda *_: None)
    return actions, freed, emptied


def test_deletes_the_misfiled_name_and_keeps_the_document(tmp_path, linked_pair):
    good, bad = linked_pair
    actions, freed, emptied = run(workbook(tmp_path, [(bad, "MISFILED", "")]))
    assert actions[0]["status"] == "deleted"
    assert not bad.exists()
    assert good.exists() and good.read_bytes() == BODY
    assert emptied == []


@pytest.mark.parametrize("decision", ["CONTAINER", "UNSURE", ""])
def test_unconfirmed_rows_are_left_alone(tmp_path, linked_pair, decision):
    _good, bad = linked_pair
    actions, freed, _ = run(workbook(tmp_path, [(bad, decision, "")]))
    assert actions == [] and freed == 0
    assert bad.exists()


def test_dry_run_deletes_nothing(tmp_path, linked_pair):
    _good, bad = linked_pair
    actions, freed, _ = run(workbook(tmp_path, [(bad, "MISFILED", "")]),
                            apply=False)
    assert actions[0]["status"] == "would delete"
    assert freed > 0 and bad.exists()


def test_refuses_to_delete_the_last_name_for_a_document(tmp_path):
    """Five shark filenames pointing at one Sci-Hub advert are all misfiled,
    but deleting all five destroys the document rather than fixing a name."""
    paths = []
    first = tmp_path / "Zhu.2011.Multiple Shark Ig H Chain Genes.pdf"
    first.write_bytes(BODY)
    paths.append(first)
    for n in ("Zhang.2013.Shark IgW C Region.pdf", "Wen.2014.Biomimetic shark skin.pdf"):
        p = tmp_path / n
        os.link(first, p)
        paths.append(p)

    book = workbook(tmp_path, [(p, "MISFILED", "") for p in paths])
    actions, freed, emptied = run(book)
    assert all(a["status"].startswith("refused") for a in actions)
    assert all(p.exists() for p in paths)
    assert freed == 0
    assert len(emptied) == 1 and emptied[0]["remaining"] == 0


def test_content_loss_proceeds_only_when_explicitly_allowed(tmp_path):
    first = tmp_path / "A.2011.One.pdf"
    first.write_bytes(BODY)
    second = tmp_path / "B.2013.Two.pdf"
    os.link(first, second)
    book = workbook(tmp_path, [(first, "MISFILED", ""), (second, "MISFILED", "")])
    actions, freed, emptied = run(book, allow_loss=True)
    assert all(a["status"] == "deleted" for a in actions)
    assert not first.exists() and not second.exists()
    assert len(emptied) == 1


def test_partial_deletion_of_a_group_is_allowed(tmp_path):
    """Deleting 18 of 19 names leaves the document reachable, so it is a
    name fix, not content loss."""
    first = tmp_path / "A.2021.One.pdf"
    first.write_bytes(BODY)
    others = []
    for n in range(3):
        p = tmp_path / f"B{n}.2021.Other.pdf"
        os.link(first, p)
        others.append(p)
    book = workbook(tmp_path, [(p, "MISFILED", "") for p in others])
    actions, _freed, emptied = run(book)
    assert all(a["status"] == "deleted" for a in actions)
    assert first.exists()
    assert emptied == []


def test_refuses_a_file_that_changed_since_the_review(tmp_path, linked_pair):
    good, bad = linked_pair
    bad.unlink()
    bad.write_bytes(b"%PDF-1.4 freshly OCRed, no longer the reviewed file")
    # Rewriting broke the link, so give the new content a sibling: otherwise
    # the content-loss gate fires first and this would not test the hash gate.
    os.link(bad, tmp_path / "sibling.pdf")
    book = workbook(tmp_path, [(bad, "MISFILED", "")])
    # Stale hash: the workbook was built before the rewrite.
    wb = openpyxl.load_workbook(book)
    ws = wb["misfiled"]
    ws.cell(2, COLS.index("sha256") + 1).value = "0" * 12
    wb.save(book)
    actions, freed, _ = run(book)
    assert actions[0]["status"] == "skipped: content changed since review"
    assert bad.exists() and freed == 0


def test_orphan_list_records_what_still_needs_fixing(tmp_path, linked_pair):
    """Deleting the file is half the repair; the record's extracted columns
    still describe the wrong paper."""
    good, bad = linked_pair
    book = workbook(tmp_path, [(bad, "MISFILED", str(good))])
    actions, _f, _e = run(book)
    out = tmp_path / "orphans.csv"
    assert app.write_orphaned_records(actions, out) == 1
    body = out.read_text()
    assert "Birth of guitarfish" in body
    assert "yes" in body            # a correct copy exists for this one

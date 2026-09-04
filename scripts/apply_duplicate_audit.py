#!/usr/bin/env python3
"""Apply the reviewed decisions from the duplicate PDF audit workbook.

Reads the ``twins`` tab of ``outputs/duplicate_pdf_audit_*.xlsx`` and deletes
the redundant filename on each approved row:

    OK         delete ``delete_name``, keep ``keep_name``
    SWAP       the other way round
    KEEP BOTH  leave both alone
    (blank)    not reviewed, leave both alone

Deletion is irreversible and these files are synced, so every row is gated on
re-hashing both paths at the moment of deletion and confirming they are still
byte-identical.  A file that changed since the scan, or whose partner has
gone, is skipped and reported rather than acted on.

Usage
-----
    python3 scripts/apply_duplicate_audit.py --workbook outputs/x.xlsx
    python3 scripts/apply_duplicate_audit.py --workbook outputs/x.xlsx --apply
"""

from __future__ import annotations

import argparse
import json
import sys
import warnings
from collections import Counter
from datetime import datetime
from pathlib import Path

warnings.filterwarnings("ignore", message=".*Data Validation.*")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dedupe_hardlink import sha256  # noqa: E402

VALID = {"OK", "SWAP", "KEEP BOTH", ""}


def read_decisions(workbook: Path) -> list[dict]:
    import openpyxl
    ws = openpyxl.load_workbook(workbook, data_only=True)["twins"]
    header = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, r))
        row["decision"] = (row.get("decision") or "").strip().upper()
        rows.append(row)
    return rows


def plan(rows: list[dict], log=print) -> tuple[list[dict], Counter]:
    counts = Counter()
    actions = []
    for row in rows:
        decision = row["decision"]
        counts[decision or "(blank)"] += 1
        if decision not in VALID:
            log(f"  unrecognised decision {decision!r} on group "
                f"{row['group_id']}, skipping")
            counts["invalid"] += 1
            continue
        if decision == "OK":
            victim, survivor = row["delete_path"], row["keep_path"]
        elif decision == "SWAP":
            victim, survivor = row["keep_path"], row["delete_path"]
        else:
            continue
        actions.append({
            "group_id": row["group_id"],
            "decision": decision,
            "delete": victim,
            "keep": survivor,
            "size_mb": row["size_mb"],
            "status": "planned",
        })
    return actions, counts


def execute(actions: list[dict], apply: bool, log=print) -> int:
    """Delete each approved file, but only after proving the survivor holds
    identical bytes right now."""
    freed = 0
    for a in actions:
        victim, survivor = Path(a["delete"]), Path(a["keep"])
        if not victim.exists():
            a["status"] = "skipped: already gone"
            continue
        if not survivor.exists():
            a["status"] = "skipped: survivor missing"
            log(f"  SKIP {victim.name}: survivor {survivor} does not exist")
            continue
        try:
            size = victim.stat().st_size
            if survivor.stat().st_size != size:
                a["status"] = "skipped: sizes now differ"
                log(f"  SKIP {victim.name}: sizes no longer match")
                continue
            if sha256(victim) != sha256(survivor):
                a["status"] = "skipped: content now differs"
                log(f"  SKIP {victim.name}: content no longer identical")
                continue
        except OSError as exc:
            a["status"] = f"skipped: {exc}"
            continue

        if not apply:
            a["status"] = "would delete"
            freed += size
            continue
        try:
            victim.unlink()
            a["status"] = "deleted"
            freed += size
        except OSError as exc:
            a["status"] = f"failed: {exc}"
            log(f"  FAILED {victim}: {exc}")
    return freed


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", type=Path, required=True)
    ap.add_argument("--apply", action="store_true",
                    help="actually delete; without this it is a dry run")
    ap.add_argument("--manifest", type=Path)
    args = ap.parse_args()

    rows = read_decisions(args.workbook)
    print(f"Read {len(rows)} reviewed rows from {args.workbook}")
    actions, counts = plan(rows)
    for k, v in sorted(counts.items()):
        print(f"  {k}: {v}")

    if not actions:
        print("\nNothing approved for deletion.")
        return 0

    print(f"\n{'Deleting' if args.apply else 'Dry run over'} {len(actions)} files...")
    freed = execute(actions, args.apply)
    done = sum(1 for a in actions if a["status"] in ("deleted", "would delete"))
    skipped = [a for a in actions if a["status"].startswith("skipped")]
    failed = [a for a in actions if a["status"].startswith("failed")]
    verb = "deleted" if args.apply else "would delete"
    print(f"  {verb} {done} files, {freed / 2 ** 30:.2f} GiB")
    print(f"  {len(skipped)} skipped, {len(failed)} failed")

    if args.apply:
        manifest = args.manifest or (
            Path("outputs") / f"duplicate_audit_applied_"
                              f"{datetime.now():%Y%m%d_%H%M%S}.json")
        manifest.parent.mkdir(parents=True, exist_ok=True)
        manifest.write_text(json.dumps(
            {"workbook": str(args.workbook),
             "generated": datetime.now().isoformat(timespec="seconds"),
             "freed_bytes": freed, "actions": actions}, indent=1))
        print(f"  manifest: {manifest}")

    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())

#!/usr/bin/env python3
"""
Feed adjudication decisions back into scripts/rag/eval/questions.jsonl.

    python3 scripts/rag/eval/apply_adjudication.py \
        outputs/rag_eval/2026-10-07_question_adjudication.xlsx --judge SD [--dry-run]

judgements tab, `relevant` column (blank rows are ignored and served again):
  Y        -> literature_id added to relevant_ids (removed from partial_ids)
  partial  -> added to partial_ids (removed from relevant_ids)
  N        -> removed from relevant_ids and partial_ids
The decision applies to `literature_id` and to every id in `also_ids`
(records with identical indexed text).
Every decision is also stored in the question's `judgements`
{literature_id: {"label", "judge", "date", "notes"}}, keyed on the stable
(question id, literature_id) pair, so build_adjudication.py never serves it
again. A question's status becomes "adjudicated" once every silver
relevant_id has a human decision and at least one row was judged.

questions tab: `new_wording` replaces the question text (status reset to
"needs_judgement", since old judgements may no longer fit, and the old
wording is kept in `previous_wording`); `drop` = Y sets status "dropped"
(run_eval.py and build_adjudication.py skip dropped questions).

The file is rewritten atomically; a copy of the previous version is kept as
questions.jsonl.bak.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import shutil
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
QUESTIONS = HERE / "questions.jsonl"
VALID = {"y": "Y", "n": "N", "partial": "partial", "p": "partial"}


def clean_id(s) -> str:
    s = str(s).strip()
    return s[:-2] if s.endswith(".0") else s


def read_sheet(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:]]


def apply(questions: list[dict], judgements: list[dict], qedits: list[dict],
          judge: str, date: str) -> dict:
    """Mutates `questions` in place; returns counts. Pure apart from that."""
    by_id = {q["id"]: q for q in questions}
    counts = {"Y": 0, "N": 0, "partial": 0, "blank": 0, "invalid": 0, "unknown_question": 0,
              "reworded": 0, "dropped": 0}
    touched = set()
    for r in judgements:
        qid = r.get("question_id")
        raw = r.get("relevant")
        if raw is None or str(raw).strip() == "":
            counts["blank"] += 1
            continue
        label = VALID.get(str(raw).strip().lower())
        if label is None:
            counts["invalid"] += 1
            continue
        q = by_id.get(qid)
        if q is None:
            counts["unknown_question"] += 1
            continue
        ids = [clean_id(r.get("literature_id"))]
        ids += [clean_id(x) for x in str(r.get("also_ids") or "").split(",") if x.strip()]
        rel = [clean_id(x) for x in q.get("relevant_ids") or []]
        part = [clean_id(x) for x in q.get("partial_ids") or []]
        rel = [x for x in rel if x not in ids]
        part = [x for x in part if x not in ids]
        if label == "Y":
            rel.extend(ids)
        elif label == "partial":
            part.extend(ids)
        q["relevant_ids"] = rel
        if part or "partial_ids" in q:
            q["partial_ids"] = part
        for lid in ids:
            q.setdefault("judgements", {})[lid] = {
                "label": label, "judge": judge, "date": date,
                "notes": (str(r["notes"]).strip() if r.get("notes") else None)}
        counts[label] += 1
        touched.add(qid)

    for r in qedits:
        q = by_id.get(r.get("id"))
        if q is None:
            continue
        if str(r.get("drop") or "").strip().lower() in ("y", "yes"):
            q["status"] = "dropped"
            counts["dropped"] += 1
            continue
        new = str(r.get("new_wording") or "").strip()
        if new and new != q["question"]:
            q["previous_wording"] = q["question"]
            q["question"] = new
            q["status"] = "needs_judgement"
            counts["reworded"] += 1
            touched.discard(q["id"])

    for qid in touched:
        q = by_id[qid]
        if q.get("status") == "dropped":
            continue
        judged = set(q.get("judgements") or {})
        silver_left = {clean_id(x) for x in q.get("relevant_ids") or []} - judged
        if not silver_left:
            q["status"] = "adjudicated"
    return counts


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--judge", required=True, help="initials of the person who judged")
    ap.add_argument("--questions", default=str(QUESTIONS))
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    wb = load_workbook(args.workbook, data_only=True)
    judgements = read_sheet(wb["judgements"]) if "judgements" in wb.sheetnames else []
    qedits = read_sheet(wb["questions"]) if "questions" in wb.sheetnames else []
    qpath = Path(args.questions)
    questions = [json.loads(l) for l in qpath.read_text(encoding="utf-8").splitlines() if l.strip()]
    counts = apply(questions, judgements, qedits, args.judge, dt.date.today().isoformat())
    print(json.dumps(counts))
    if args.dry_run:
        print("dry run: questions.jsonl not written")
        return
    shutil.copy2(qpath, qpath.with_suffix(".jsonl.bak"))
    tmp = qpath.with_suffix(".jsonl.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        for q in questions:
            f.write(json.dumps(q, ensure_ascii=False) + "\n")
    os.replace(tmp, qpath)
    print(f"wrote {qpath}")


if __name__ == "__main__":
    main()
